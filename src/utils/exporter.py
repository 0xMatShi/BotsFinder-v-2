import os
from openpyxl import Workbook
from openpyxl.styles import Font


OUTPUT_FOLDER = "results"

def get_unique_filepath(folder, base_name="wallets", extension="xlsx"):
    """
    Генерирует уникальный путь к файлу внутри указанной папки.
    Пример: results/wallets_1.xlsx, results/wallets_2.xlsx ...
    """
    counter = 1
    while True:
        filename = f"{base_name}_{counter}.{extension}"
        full_path = os.path.join(folder, filename)
        
        # Если такого файла в папке нет, возвращаем этот путь
        if not os.path.exists(full_path):
            return full_path
        counter += 1

def save_wallets_to_xlsx(wallets_data):
    """
    Сохраняет данные в Excel с гиперссылками в папку results.
    """
    if not wallets_data:
        print("Нет данных для сохранения.")
        return

    # 1. Создаем папку results, если её нет
    if not os.path.exists(OUTPUT_FOLDER):
        os.makedirs(OUTPUT_FOLDER)

    # 2. Генерируем уникальный путь файла внутри папки
    filepath = get_unique_filepath(OUTPUT_FOLDER)

    # 3. Создаем книгу и лист
    wb = Workbook()
    ws = wb.active
    ws.title = "Whales" #type: ignore

    # 4. Заголовки
    headers = ["Wallet", "Predictions", "Polymarket", "Hashdive"]
    ws.append(headers) #type: ignore

    # Стиль для заголовков (жирный)
    for cell in ws[1]: #type: ignore
        cell.font = Font(bold=True)

    # Стиль для ссылок (синий, подчеркнутый)
    link_font = Font(color="0563C1", underline="single")

    # 5. Заполняем данными
    for item in wallets_data:
        wallet = item['address']
        traded_count = item['traded']
        
        poly_url = f"https://polymarket.com/@{wallet}"
        hash_url = f"https://hashdive.com/Analyze_User?user_address={wallet}"

        # Добавляем строку
        ws.append([wallet, traded_count, poly_url, hash_url]) #type: ignore

        # Получаем номер текущей строки
        current_row = ws.max_row #type: ignore

        # --- Настраиваем ссылку Polymarket (Колонка C = 3) ---
        cell_poly = ws.cell(row=current_row, column=3) #type: ignore
        cell_poly.hyperlink = poly_url #type: ignore
        cell_poly.value = "Polymarket"  # Текст ссылки #type: ignore
        cell_poly.font = link_font

        # --- Настраиваем ссылку Hashdive (Колонка D = 4) ---
        cell_hash = ws.cell(row=current_row, column=4) #type: ignore
        cell_hash.hyperlink = hash_url #type: ignore
        cell_hash.value = "Hashdive" #type: ignore
        cell_hash.font = link_font

    # 6. Настройка ширины колонок
    ws.column_dimensions['A'].width = 45  # Wallet #type: ignore
    ws.column_dimensions['B'].width = 20  # Predictions #type: ignore
    ws.column_dimensions['C'].width = 15  # Polymarket (текст стал короче) #type: ignore
    ws.column_dimensions['D'].width = 15  # Hashdive (текст стал короче) #type: ignore

    # 7. Сохранение
    try:
        wb.save(filepath)
        # Получаем полный абсолютный путь для вывода в консоль
        full_abs_path = os.path.abspath(filepath)
        print(f"\n💾 Таблица успешно сохранена: {filepath}")
        print(f"   Полный путь: {full_abs_path}")
    except PermissionError:
        print(f"\n❌ Ошибка: Файл {filepath} открыт в Excel. Закройте его и попробуйте снова.")
    except Exception as e:
        print(f"\n❌ Ошибка при сохранении xlsx: {e}")